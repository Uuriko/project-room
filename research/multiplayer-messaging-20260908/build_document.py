from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

HERE = Path(__file__).resolve().parent
source = (HERE / 'report-source.md').read_text()
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for border in doc.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.68)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
for name, size in [('Heading 1', 16), ('Heading 2', 12)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(15 if size == 16 else 9)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = 'PROJECT ROOM'
header.runs[0].font.name = 'Arial'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = sec.footer.paragraphs[0]
footer.paragraph_format.space_before = Pt(0)
footer.alignment = 2
r = footer.add_run('September 2026   •   ')
r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Project Room Collaboration and Messaging Plan'
doc.core_properties.subject = 'Research and integration proposal'
doc.core_properties.author = 'Codex for Project Room'

def text_runs(p, text):
    cursor = 0
    for m in re.finditer(r'\[([^\]]+)\]\((https?://[^\s]+?)\)', text):
        p.add_run(text[cursor:m.start()])
        link = OxmlElement('w:hyperlink')
        rid = p.part.relate_to(m.group(2), 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        link.set(qn('r:id'), rid)
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '333333'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props); t = OxmlElement('w:t'); t.text = m.group(1); run.append(t)
        link.append(run); p._p.append(link)
        cursor = m.end()
    p.add_run(text[cursor:])

def table(rows):
    columns = len(rows[0])
    t = doc.add_table(rows=0, cols=columns)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [1.32, 2.06, 3.52] if columns == 3 else [2.6, 4.3]
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + side); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
    props.append(borders)
    for idx, values in enumerate(rows):
        row = t.add_row()
        trp = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trp.append(no_split)
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); trp.append(repeat)
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '90'); el.set(qn('w:type'), 'dxa'); margin.append(el)
            cp.append(margin)
            if idx == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAEAEA'); cp.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            text_runs(p, value)
            for r in p.runs:
                r.font.size = Pt(10)
                r.bold = idx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

lines = source.splitlines()
i = 0
refs = False
while i < len(lines):
    line = lines[i].strip()
    i += 1
    if not line: continue
    if line.startswith('|'):
        entries = [line]
        while i < len(lines) and lines[i].strip().startswith('|'):
            entries.append(lines[i].strip()); i += 1
        rows = [[v.strip() for v in x.strip('|').split('|')] for x in entries if not re.match(r'^\|[\s:|\-]+\|$', x)]
        table(rows); continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title'); continue
    if line.startswith('## '):
        title = line[3:]
        if title == 'Source notes':
            doc.add_page_break(); refs = True
        doc.add_heading(title, 1); continue
    if line.startswith('### '):
        doc.add_heading(line[4:], 2); continue
    if line == 'Research and integration proposal':
        doc.add_paragraph(line, 'Subtitle'); continue
    style = 'Normal'
    if line.startswith('- '): line, style = line[2:], 'List Bullet'
    elif re.match(r'^\d+\. ', line):
        # Preserve the source's intentional restart instead of continuing Word's shared list.
        style = 'Normal'
    p = doc.add_paragraph(style=style)
    if re.match(r'^\d+\. ', line):
        p.paragraph_format.space_after = Pt(0)
    text_runs(p, line)
    if refs:
        p.paragraph_format.line_spacing = 1.06
        p.paragraph_format.space_after = Pt(7)
        for run in p.runs: run.font.size = Pt(9.5)

out = HERE / 'Project Room Collaboration and Messaging Plan.docx'
doc.save(out)
print(str(out))
print(f'Source words: {len(source.split())}; paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}')
